"""PC-side logger for the STM32 slip/load test rig.

Connects to the Nucleo's ST-LINK virtual COM port and saves every motor run
automatically - no TEST/STOP needed. At startup it asks three questions: what
equipment is on the hook, how much it weighs, and what frequency you intend to
test it at. Those three answers name the session. Each time the motor runs,
its steady-state averages are kept in memory; when the motor stops, the logger
asks for the RPM shown on the SERVO display and adds one row to the summary
workbook.

  data/2026-10-07_Metalhousing_baseweight-1.46kg_testedfreq-6.36hz/
    2026-10-07_Metalhousing_baseweight-1.46kg_testedfreq-6.36hz_summary.xlsx
    2026-10-07_Metalhousing_baseweight-1.46kg_testedfreq-6.36hz_session_log.txt

Only these two files are written - no per-run CSV and no all-rows file - so a
session folder stays short and the summary opens directly in Excel, already
editable (no CSV import step).

Commands typed in the window:
  291 / 291 down     after a run: the SERVO RPM (and optionally the direction)
  Enter              after a run: skip the RPM
  WEIGHT 2.5         change the equipment weight for the following runs
  QUIT               save and exit
Anything else (VFDCHECK, DEBUG 1, ...) is sent to the board.

Usage:
  python vfd_logger.py              # auto-detects the ST-LINK COM port
  python vfd_logger.py --port COM5  # or name it yourself

Type QUIT (or press Ctrl+C) to exit.
"""

import argparse
import datetime
import os
import re
import sys
import threading
import time
import traceback

try:
    import serial
    import serial.tools.list_ports
except ImportError:
    sys.exit("pyserial is not installed. Run:  pip install pyserial")

try:
    from openpyxl import Workbook
except ImportError:
    sys.exit("openpyxl is not installed. Run:  pip install openpyxl")

# Matches the firmware's CSV header; replaced by the header line the board prints
# at startup ("# millis,freq_hz,...") if the logger sees it.
DEFAULT_COLUMNS = [
    "millis", "freq_hz", "freq_source", "sync_rpm", "actual_rpm", "slip",
    "current_a", "current_source", "known_load_kg", "test_id", "steady",
]
# One row per run. "weight_kg" and "target_freq_hz" repeat the session's own
# identity (also in the file name) so a row still makes sense if this sheet is
# ever copied out on its own; everything else is measured, not typed in twice.
SUMMARY_COLUMNS = [
    "run", "direction", "weight_kg", "target_freq_hz", "start_time", "run_s",
    "steady_rows", "avg_freq_hz", "avg_current_a", "sync_rpm", "actual_rpm", "slip",
]
STLINK_VID = 0x0483
RUNNING_FREQ_HZ = 0.5  # output frequency above this counts as "motor running"
NO_REPLY_WARNING_S = 2.0

# Steady rows more than this below the highest steady speed in a run (pauses while
# turning the pot, slowing down) are left out of that run's averages.
PLATEAU_TOL_HZ = 0.2
RPM_ANSWER_RE = re.compile(r"^(?:RPM\s*)?(\d+(?:\.\d+)?)(?:\s+(UP|DOWN|U|D))?$")


def find_stlink_port():
    ports = list(serial.tools.list_ports.comports())
    for p in ports:
        if p.vid == STLINK_VID or "STLink" in (p.description or "") or "STM" in (p.description or ""):
            return p.device
    names = ", ".join(f"{p.device} ({p.description})" for p in ports) or "none found"
    sys.exit(f"Couldn't find the ST-LINK COM port. Ports seen: {names}\n"
             f"Run again with --port COMx (check Device Manager -> Ports).")


def now_str():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]


def report_crash(context):
    # Prints the error and saves it to crash_log.txt next to the script, so a problem
    # can be diagnosed instead of the window just closing.
    details = traceback.format_exc()
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "crash_log.txt")
    try:
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"--- {now_str()} {context}\n{details}\n")
    except OSError:
        pass
    print(f"\n!! ERROR {context}\n{details}!! saved to {path} - send this to Claude\n")


def kg_label(kg):
    return f"{float(kg):g}kg"


def hz_label(hz):
    return f"{float(hz):g}hz"


def safe_name(text):
    # Strips spaces and anything that isn't safe in a file/folder name, so
    # "Metal housing" -> "Metalhousing" instead of needing manual cleanup later.
    cleaned = re.sub(r"[^A-Za-z0-9_-]", "", text)
    return cleaned or "equipment"


def ask_text(question):
    while True:
        try:
            text = input(f">> {question}: ").strip()
        except EOFError:
            return "equipment"
        if text:
            return text
        print(">> type something, e.g. Metal housing")


def ask_kg(question):
    while True:
        try:
            text = input(f">> {question} (kg): ").strip()
        except EOFError:
            return 0.0
        try:
            kg = float(text)
            if kg >= 0:
                return kg
        except ValueError:
            pass
        print(">> enter a number, e.g. 1.46")


def ask_hz(question):
    while True:
        try:
            text = input(f">> {question} (Hz): ").strip()
        except EOFError:
            return 0.0
        try:
            hz = float(text)
            if hz >= 0:
                return hz
        except ValueError:
            pass
        print(">> enter a number, e.g. 6.36")


class Logger:
    def __init__(self, ser, out_dir, max_run_s, show_all, equipment_name, weight_kg,
                 target_freq_hz, session_stamp):
        self.ser = ser
        self.out_dir = out_dir
        self.max_run_s = max_run_s
        self.show_all = show_all
        self.equipment_name = equipment_name
        self.weight_kg = weight_kg
        self.target_freq_hz = target_freq_hz
        self.columns = list(DEFAULT_COLUMNS)
        self.lock = threading.Lock()
        self.running = True
        self.closed = False
        self.last_reply_time = 0.0

        self.run_count = 0
        self.run = None        # the run in progress: rows, start time
        self.pending = None    # finished run waiting for its SERVO RPM

        prefix = f"{session_stamp}_"
        self.session_log = open(os.path.join(out_dir, f"{prefix}session_log.txt"), "a", encoding="utf-8")
        self.session_log.write(
            f"{now_str()} equipment={equipment_name} weight_kg={weight_kg:g} "
            f"target_freq_hz={target_freq_hz:g}\n"
        )

        self.summary_path = os.path.join(out_dir, f"{prefix}summary.xlsx")
        self.summary_name = os.path.basename(self.summary_path)
        if os.path.exists(self.summary_path):
            # Re-opening an existing session (e.g. after a crash): append, don't overwrite.
            from openpyxl import load_workbook
            self.wb = load_workbook(self.summary_path)
            self.ws = self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Summary"
            self.ws.append(SUMMARY_COLUMNS)
            self.wb.save(self.summary_path)

        # Column number-formats for the numeric fields written by write_summary(),
        # keyed by SUMMARY_COLUMNS index (0-based) so Excel shows sensible decimals
        # instead of raw floats.
        self._col_formats = {
            SUMMARY_COLUMNS.index("weight_kg"): "0.00",
            SUMMARY_COLUMNS.index("target_freq_hz"): "0.00",
            SUMMARY_COLUMNS.index("run_s"): "0.0",
            SUMMARY_COLUMNS.index("avg_freq_hz"): "0.00",
            SUMMARY_COLUMNS.index("avg_current_a"): "0.000",
            SUMMARY_COLUMNS.index("sync_rpm"): "0.0",
            SUMMARY_COLUMNS.index("actual_rpm"): "0.0",
            SUMMARY_COLUMNS.index("slip"): "0.0000",
        }

    # --- receiving -------------------------------------------------------------
    def reader_loop(self):
        buf = b""
        while self.running:
            try:
                chunk = self.ser.read(256)
            except serial.SerialException as e:
                print(f"\n!! serial port error: {e}")
                self.running = False
                break
            if not chunk:
                continue
            buf += chunk
            while b"\n" in buf:
                raw, buf = buf.split(b"\n", 1)
                line = raw.decode("utf-8", errors="replace").strip()
                if line:
                    try:
                        self.handle_line(line)
                    except Exception:
                        report_crash(f"while handling the line: {line!r}")

    def handle_line(self, line):
        with self.lock:
            if self.closed:
                return
            stamp = now_str()
            self.session_log.write(f"{stamp} < {line}\n")
            self.session_log.flush()

            if line.startswith("#"):
                print(line)
                self.last_reply_time = time.monotonic()
                if line.startswith("# millis,"):
                    cols = line[2:].split(",")
                    if cols != self.columns:
                        self.columns = cols
                return

            fields = line.split(",")
            if len(fields) != len(self.columns):
                return  # partial/garbled line - kept in the session log only
            row = dict(zip(self.columns, fields))
            running = self.is_running(row)
            # Idle rows aren't shown, so the screen stays still while you type.
            if self.show_all or running or self.run is not None:
                print(line)

            if running and self.run is None:
                self.start_run()
            if self.run is not None:
                self.run["rows"].append(row)
                elapsed = time.monotonic() - self.run["t0"]
                if running and not self.run["warned"] and elapsed >= self.max_run_s:
                    self.run["warned"] = True
                    print(f"\a!!! {elapsed:.1f} s running - RELEASE THE BUTTON (--max-run {self.max_run_s:g})")
                if not running:
                    self.end_run()

    @staticmethod
    def is_running(row):
        try:
            return float(row.get("freq_hz", "0")) > RUNNING_FREQ_HZ
        except ValueError:
            return False

    # --- runs ----------------------------------------------------------------------
    def start_run(self):
        if self.pending is not None:
            print(f">> run {self.pending['n']}: no RPM entered - saved without it")
            self.write_summary(self.pending, None, "")
        self.run_count += 1
        self.run = {"n": self.run_count, "start_time": now_str(), "weight": self.weight_kg,
                    "rows": [], "t0": time.monotonic(), "warned": False}

    def end_run(self):
        run = self.run
        self.run = None
        run["seconds"] = time.monotonic() - run["t0"]
        run.update(self.steady_stats(run["rows"]))
        self.pending = run
        if run["steady_rows"]:
            print(f">> run {run['n']} saved ({run['seconds']:.1f} s): {run['avg_freq_hz']:.2f} Hz, "
                  f"{run['avg_current_a']:.3f} A over {run['steady_rows']} steady rows")
        else:
            print(f">> run {run['n']} saved ({run['seconds']:.1f} s) - too short, no steady speed")
        print(f">> type the SERVO RPM for run {run['n']} (e.g. 291, or 291 down) and press Enter,"
              f" or just Enter to skip")

    @staticmethod
    def steady_stats(rows):
        # Average only the rows at the highest constant speed the run held.
        steady = []
        for r in rows:
            try:
                if r.get("steady") == "1":
                    steady.append((float(r["freq_hz"]), float(r["current_a"]), float(r["sync_rpm"])))
            except (KeyError, ValueError):
                pass
        if not steady:
            return {"steady_rows": 0}
        top = max(f for f, _, _ in steady)
        plateau = [s for s in steady if s[0] >= top - PLATEAU_TOL_HZ]
        n = len(plateau)
        return {
            "steady_rows": n,
            "avg_freq_hz": sum(s[0] for s in plateau) / n,
            "avg_current_a": sum(s[1] for s in plateau) / n,
            "sync_rpm": sum(s[2] for s in plateau) / n,
        }

    def write_summary(self, run, rpm, direction):
        slip = None
        if rpm is not None and run.get("sync_rpm"):
            slip = (run["sync_rpm"] - rpm) / run["sync_rpm"]
        row_values = [
            run["n"], direction, run["weight"], self.target_freq_hz, run["start_time"],
            run["seconds"], run["steady_rows"], run.get("avg_freq_hz"), run.get("avg_current_a"),
            run.get("sync_rpm"), rpm, slip,
        ]
        self.ws.append(row_values)
        row_idx = self.ws.max_row
        for col_idx, number_format in self._col_formats.items():
            self.ws.cell(row=row_idx, column=col_idx + 1).number_format = number_format
        self.wb.save(self.summary_path)
        return f"{slip:.4f}" if slip is not None else ""

    # --- typed input -----------------------------------------------------------
    def handle_input(self, text):
        cmd = text.strip().upper()
        with self.lock:
            if self.pending is not None:
                if not cmd:
                    self.write_summary(self.pending, None, "")
                    print(f">> run {self.pending['n']} added to {self.summary_name} without RPM")
                    self.pending = None
                    return
                m = RPM_ANSWER_RE.match(cmd)
                if m:
                    rpm = float(m.group(1))
                    direction = {"U": "up", "D": "down"}.get(m.group(2), (m.group(2) or "").lower())
                    slip = self.write_summary(self.pending, rpm, direction)
                    print(f">> run {self.pending['n']}: RPM {rpm:g}{' ' + direction if direction else ''}"
                          f"{', slip ' + slip if slip else ''} -> added to {self.summary_name}")
                    self.pending = None
                    return
            if not cmd:
                return
            parts = cmd.split()
            if parts[0] in ("WEIGHT", "LOAD") and len(parts) == 2:
                try:
                    kg = float(parts[1])
                except ValueError:
                    kg = -1
                if kg >= 0:
                    self.weight_kg = kg
                    self.session_log.write(f"{now_str()} weight_kg={kg:g}\n")
                    print(f">> equipment weight is now {kg:g} kg for the next runs")
                else:
                    print(">> usage: WEIGHT 2.5")
                return
        self.send(cmd)

    def send(self, cmd):
        with self.lock:
            self.session_log.write(f"{now_str()} > {cmd}\n")
            self.session_log.flush()
        sent_at = time.monotonic()
        self.ser.write((cmd + "\r\n").encode("ascii", errors="replace"))
        threading.Timer(NO_REPLY_WARNING_S, self.check_reply, args=(sent_at, cmd)).start()

    def check_reply(self, sent_at, cmd):
        if self.running and self.last_reply_time < sent_at:
            print(f"!! no reply to '{cmd}' from the board - is it running (LD2 blinking)?")

    def close(self):
        self.running = False
        with self.lock:
            self.closed = True
            if self.run is not None:
                run = self.run
                run["seconds"] = time.monotonic() - run["t0"]
                run.update(self.steady_stats(run["rows"]))
                print(f">> run {run['n']} was still going - saved without a final RPM reading")
                self.write_summary(run, None, "")
            if self.pending is not None:
                self.write_summary(self.pending, None, "")
            self.wb.save(self.summary_path)
            self.session_log.close()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--port", help="serial port, e.g. COM5 (default: auto-detect ST-LINK)")
    ap.add_argument("--baud", type=int, default=115200)
    ap.add_argument("--out", default="data", help="folder for session folders (default: ./data)")
    ap.add_argument("--max-run", type=float, default=5.0,
                    help="beep and warn when the motor has run this many seconds (default: 5.0)")
    ap.add_argument("--show-all", action="store_true",
                    help="show every data row, including idle ones (default: only while the motor runs)")
    ap.add_argument("--equipment", help="name of what's on the hook; skips the startup question")
    ap.add_argument("--weight", type=float, help="its weight in kg; skips the startup question")
    ap.add_argument("--freq", type=float, help="intended test frequency in Hz; skips the startup question")
    args = ap.parse_args()

    port = args.port or find_stlink_port()
    equipment_name = args.equipment or ask_text("Equipment being tested (e.g. Metal housing)")
    weight_kg = args.weight if args.weight is not None else ask_kg(f"Weight of {equipment_name}")
    target_freq_hz = args.freq if args.freq is not None else ask_hz("Intended test frequency")

    date_stamp = f"{datetime.datetime.now():%Y-%m-%d}"
    session_stamp = (f"{date_stamp}_{safe_name(equipment_name)}_baseweight-{kg_label(weight_kg)}"
                      f"_testedfreq-{hz_label(target_freq_hz)}")
    out_dir = os.path.join(args.out, session_stamp)
    os.makedirs(out_dir, exist_ok=True)

    try:
        ser = serial.Serial(port, args.baud, timeout=0.2)
    except serial.SerialException as e:
        sys.exit(f"Couldn't open {port}: {e}\n"
                 f"Close anything else using it (PuTTY, Tera Term) and try again.")

    log = Logger(ser, out_dir, args.max_run, args.show_all, equipment_name, weight_kg,
                 target_freq_hz, session_stamp)
    print(f">> connected to {port}: {equipment_name}, {weight_kg:g} kg, testing at {target_freq_hz:g} Hz")
    print(f">> saving to {out_dir}")
    print(">> every motor run is saved automatically - after each run, type the SERVO RPM")
    print(">> commands: WEIGHT <kg> (change equipment weight)   VFDCHECK   QUIT")
    print(f">> run-time warning at {args.max_run:g} s (change with --max-run)\n")

    reader = threading.Thread(target=log.reader_loop, daemon=True)
    reader.start()

    try:
        while log.running:
            try:
                text = input()
            except EOFError:
                break
            if text.strip().upper() in ("QUIT", "EXIT"):
                break
            try:
                log.handle_input(text)
            except Exception:
                report_crash(f"while handling {text!r}")
    except KeyboardInterrupt:
        pass
    finally:
        log.close()
        time.sleep(0.3)
        ser.close()
        print(f">> session saved in {out_dir}")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception:
        report_crash("in the logger")
        sys.exit(1)
